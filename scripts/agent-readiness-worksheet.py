"""The agent readiness worksheet: a fillable, printable A4 PDF (two pages).

Day 8 of the AI bootcamp carried four reference sections that made the article
heavy to read and were never going to be read in the browser anyway: they are
things you work through while designing an agent. They live here instead.

Uses the site's own fonts (decompressed from src/app/fonts/*.woff2 at run time)
and colour tokens, so the download matches the website and the bootcamp.

    python3 scripts/agent-readiness-worksheet.py
"""
import os, tempfile
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import Color
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "public", "downloads", "agent-readiness-worksheet.pdf")


def hexc(h, a=1.0):
    h = h.lstrip("#")
    return Color(*(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)), alpha=a)


PAPER = hexc("EDE7DA"); NAVY = hexc("0A1D2B"); ORANGE = hexc("dd3c13"); INK = hexc("111111")
MUTED = hexc("5F5E5A"); BOX = hexc("E6E0D4"); FIELD_BORDER = hexc("CFC7B7"); WHITE = hexc("FFFFFF")
ORANGE_ON_NAVY = hexc("F0A48F")


def _fonts():
    from fontTools.ttLib import TTFont as FT
    tmp = tempfile.mkdtemp(prefix="ecp-fonts-")
    for name, file in (("RC", "roboto-condensed-700"), ("IN", "inter")):
        dst = os.path.join(tmp, f"{file}.ttf")
        f = FT(os.path.join(ROOT, "src", "app", "fonts", f"{file}.woff2")); f.flavor = None; f.save(dst)
        pdfmetrics.registerFont(TTFont(name, dst))


_fonts()

ARROW = [(10, 66), (28, 84), (62, 50), (62, 75), (86, 75), (86, 10), (21, 10), (21, 34), (46, 34)]
W, H = A4
M = 36


def arrow(c, x, y, size, col):
    s = size / 96; p = c.beginPath()
    p.moveTo(x + ARROW[0][0] * s, y + size - ARROW[0][1] * s)
    for px, py in ARROW[1:]: p.lineTo(x + px * s, y + size - py * s)
    p.close(); c.setFillColor(col); c.drawPath(p, fill=1, stroke=0)


def wrap(c, text, font, size, width):
    out, line = [], ""
    for w in text.split():
        t = (line + " " + w).strip()
        if c.stringWidth(t, font, size) <= width: line = t
        else: out.append(line); line = w
    if line: out.append(line)
    return out


def field(c, name, tip, x, y, w, h, multiline=True):
    c.acroForm.textfield(name=name, tooltip=tip, x=x, y=y, width=w, height=h,
                         fieldFlags="multiline" if multiline else "", fontName="Helvetica", fontSize=8,
                         borderWidth=0.6, borderColor=FIELD_BORDER, fillColor=WHITE, textColor=INK, forceBorder=True)


def header(c, subtitle):
    c.setFillColor(NAVY); c.rect(0, H - 44, W, 44, fill=1, stroke=0)
    arrow(c, M, H - 44 + 15, 13, PAPER)
    c.setFont("RC", 12.5); c.setFillColor(PAPER); c.drawString(M + 19, H - 44 + 16.5, "european campaign")
    c.setFillColor(hexc("EDE7DA", 0.6))
    c.drawString(M + 19 + c.stringWidth("european campaign ", "RC", 12.5), H - 44 + 16.5, "playbook")
    c.setFont("IN", 7.5); c.setFillColor(hexc("EDE7DA", 0.75))
    c.drawRightString(W - M, H - 44 + 17, subtitle)


def footer(c, page):
    c.setStrokeColor(FIELD_BORDER); c.setLineWidth(0.5); c.line(M, 46, W - M, 46)
    c.setFont("IN", 7); c.setFillColor(MUTED)
    c.drawString(M, 32, "campaignplaybook.eu/articles/ai-bootcamp-eu-affairs-day-8")
    c.drawRightString(W - M, 32, f"Page {page} of 2")


QUESTIONS = [
    "What single recurring problem does it solve?",
    "What triggers it?",
    "What are the exact steps?",
    "Which sources can it trust?",
    "Which tools does it need?",
    "Which actions can it take without approval?",
    "Which actions always require approval?",
    "What must it never do?",
    "Who reviews the result?",
    "How will we know it is working?",
]

NEVER = [
    "Unsupervised external political communication",
    "Contacting stakeholders in your name",
    "Changing organisational positions",
    "Publishing social content",
    "Making sensitive strategic decisions",
    "Interpreting ambiguous political developments without review",
    "Maintaining sensitive personal profiles",
    "Taking irreversible actions",
]

CHECKLIST = [
    ("Instructions", [
        "Mission is specific.", "Role is defined.", "Audience is defined.",
        "Workflow is written step by step.", "Output format is defined.",
        "Uncertainty rules are defined."]),
    ("Knowledge", [
        "Sources are trusted.", "Old material has been reviewed.",
        "Rejected drafts are excluded.", "Important primary sources are included.",
        "AI-generated material is checked before it counts as fact."]),
    ("Capabilities", [
        "Every connected tool has a clear purpose.",
        "Permissions are no broader than necessary.",
        "External actions are limited.",
        "Internal and external actions are distinguished."]),
    ("Boundaries", [
        "Prohibited actions are explicit.", "Human approval points are defined.",
        "Sensitive situations trigger escalation.", "Failure behaviour is defined.",
        "Recipient lists are controlled."]),
    ("Testing", [
        "Workflow has been run manually.", "Outputs have been reviewed.",
        "Instructions have been revised.", "Sources have been checked.",
        "Scheduling comes only after the workflow is stable."]),
]

MISTAKES = [
    "Building an agent before documenting the existing workflow.",
    "Giving it a vague role.",
    "Treating instructions as a one-line prompt.",
    "Feeding it unreviewed AI-generated knowledge.",
    "Uploading every document available instead of selecting trustworthy material.",
    "Giving it more capabilities than the job requires.",
    "Forgetting to define what it cannot do.",
    "Automating external communication too early.",
    "Failing to name a human owner.",
    "Leaving out when the agent should stop or escalate.",
    "Scheduling a broken workflow and letting it fail repeatedly.",
    "Assuming autonomy removes the need for verification.",
]


def build():
    c = canvas.Canvas(OUT, pagesize=A4)
    c.setTitle("The agent readiness worksheet")
    c.setAuthor("european campaign playbook")

    # ── page 1: the ten questions, and what to keep it away from ───────────
    c.setFillColor(PAPER); c.rect(0, 0, W, H, fill=1, stroke=0)
    header(c, "AI bootcamp for EU affairs  ·  day 8")

    c.setFont("RC", 21); c.setFillColor(ORANGE); c.drawString(M, H - 76, "The agent readiness test")
    c.setFont("IN", 7.8); c.setFillColor(MUTED)
    c.drawString(M, H - 91, "Ten questions to answer before you let an agent run. Without clear answers, stay with a project")
    c.drawString(M, H - 101, "or a manual workflow for now. That is good system design.")

    # Meta row on its own line, clear of the title above it.
    fx = M
    for label, fw in (("Agent", 150), ("Owner", 110), ("Date", 80)):
        c.setFont("IN", 7.5); c.setFillColor(MUTED); c.drawString(fx, H - 126, label)
        lw = c.stringWidth(label, "IN", 7.5)
        field(c, label.lower(), label, fx + lw + 6, H - 130, fw, 15, multiline=False)
        fx += lw + 6 + fw + 18

    # Each question and its answer box share a row.
    qw, row_h, box_h = 150, 41, 33
    y = H - 152
    for i, q in enumerate(QUESTIONS, 1):
        c.setFont("RC", 10); c.setFillColor(ORANGE); c.drawString(M, y - 9, f"{i:02}")
        c.setFont("IN", 8.2); c.setFillColor(NAVY)
        ty = y - 9
        for ln in wrap(c, q, "IN", 8.2, qw):
            c.drawString(M + 20, ty, ln); ty -= 10.5
        field(c, f"q{i}", q, M + 186, y - box_h + 4, W - M - (M + 186), box_h)
        y -= row_h

    # The panel is sized to its own content rather than to what is left over.
    rows = (len(NEVER) + 1) // 2
    panel_h = 54 + rows * 14
    ptop = y - 10
    c.setFillColor(NAVY); c.roundRect(M, ptop - panel_h, W - 2 * M, panel_h, 4, fill=1, stroke=0)
    c.setFont("RC", 13); c.setFillColor(PAPER); c.drawString(M + 12, ptop - 22, "Where not to start")
    c.setFont("IN", 6.8); c.setFillColor(ORANGE_ON_NAVY)
    c.drawString(M + 12, ptop - 34, "THE BEST FIRST AGENT IS USUALLY BORING. A BORING WORKFLOW IS EASIER TO UNDERSTAND, TEST AND IMPROVE.")
    colw = (W - 2 * M - 24) / 2
    for i, item in enumerate(NEVER):
        col, rowi = divmod(i, rows)
        x = M + 12 + col * colw
        iy = ptop - 52 - rowi * 14
        c.setFillColor(ORANGE); c.setFont("RC", 9); c.drawString(x, iy, "x")
        c.setFillColor(hexc("EDE7DA", 0.9)); c.setFont("IN", 7.8); c.drawString(x + 12, iy, item)

    footer(c, 1)
    c.showPage()

    # ── page 2: the build checklist, and the mistakes to check against ─────
    c.setFillColor(PAPER); c.rect(0, 0, W, H, fill=1, stroke=0)
    header(c, "AI bootcamp for EU affairs  ·  day 8")

    c.setFont("RC", 21); c.setFillColor(ORANGE); c.drawString(M, H - 74, "The build checklist")
    c.setFont("IN", 7.5); c.setFillColor(MUTED)
    c.drawString(M, H - 87, "Tick each line before the agent runs on a schedule. Anything unticked is a decision you have not made yet.")

    y = H - 108
    colw = (W - 2 * M - 16) / 2
    col_x = [M, M + colw + 16]
    col_y = [y, y]
    n = 0
    for title, items in CHECKLIST:
        col = 0 if n < 3 else 1
        x, yy = col_x[col], col_y[col]
        box_h = 26 + len(items) * 14
        c.setFillColor(BOX); c.roundRect(x, yy - box_h, colw, box_h, 4, fill=1, stroke=0)
        c.setFont("RC", 11.5); c.setFillColor(NAVY); c.drawString(x + 10, yy - 17, title)
        iy = yy - 33
        for j, item in enumerate(items):
            c.acroForm.checkbox(name=f"{title.lower()}_{j}", tooltip=item, x=x + 10, y=iy - 2.5, size=9,
                                buttonStyle="check", borderColor=NAVY, fillColor=WHITE, textColor=NAVY,
                                forceBorder=True, borderWidth=0.7)
            c.setFont("IN", 7.6); c.setFillColor(INK)
            for ln in wrap(c, item, "IN", 7.6, colw - 36)[:1]:
                c.drawString(x + 24, iy, ln)
            iy -= 14
        col_y[col] = yy - box_h - 12
        n += 1

    ytop = min(col_y) - 4
    c.setFont("RC", 13); c.setFillColor(NAVY); c.drawString(M, ytop, "Twelve mistakes to check yourself against")
    yy = ytop - 16
    for i, m in enumerate(MISTAKES):
        x = M if i < 6 else M + colw + 16
        if i == 6: yy = ytop - 16
        c.setFillColor(ORANGE); c.setFont("IN", 7.8); c.drawString(x, yy, "•")
        c.setFillColor(INK); c.setFont("IN", 7.6)
        for ln in wrap(c, m, "IN", 7.6, colw - 14):
            c.drawString(x + 9, yy, ln); yy -= 9.6
        yy -= 3.5

    # Somewhere to write while working through the list above.
    ny = min(yy, 230)
    c.setFont("RC", 13); c.setFillColor(NAVY); c.drawString(M, ny, "What still has to be decided")
    field(c, "notes", "What still has to be decided", M, 78, W - 2 * M, ny - 92)

    c.setFont("RC", 10); c.setFillColor(ORANGE)
    c.drawCentredString(W / 2, 60, "You design the system. The agent runs it.")
    footer(c, 2)
    c.save()




def build_docx():
    """The same worksheet as an editable Word file, for people who would
    rather type into it than use a PDF form."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    def heading(text, size=18, colour="dd3c13", space_before=14):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(colour.upper())
        return p

    def muted(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string("5F5E5A")

    title = doc.add_paragraph()
    tr = title.add_run("The agent readiness worksheet")
    tr.bold = True
    tr.font.size = Pt(24)
    tr.font.color.rgb = RGBColor.from_string("0A1D2B")
    muted("AI bootcamp for EU affairs, day 8  ·  european campaign playbook  ·  campaignplaybook.eu")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("Agent", "Owner", "Date")):
        cell.text = f"{label}: "

    heading("The agent readiness test")
    muted("Ten questions to answer before you let an agent run. Without clear answers, stay with a project or a manual workflow for now. That is good system design.")
    qt = doc.add_table(rows=len(QUESTIONS) + 1, cols=2)
    qt.style = "Table Grid"
    qt.rows[0].cells[0].paragraphs[0].add_run("Question").bold = True
    qt.rows[0].cells[1].paragraphs[0].add_run("Your answer").bold = True
    for i, q in enumerate(QUESTIONS, 1):
        qt.rows[i].cells[0].text = f"{i}. {q}"
        qt.rows[i].cells[1].text = ""

    heading("Where not to start", size=15, colour="0A1D2B")
    muted("The best first agent is usually boring. A boring workflow is easier to understand, test and improve.")
    for item in NEVER:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()

    heading("The build checklist")
    muted("Tick each line before the agent runs on a schedule. Anything unticked is a decision you have not made yet.")
    for group, items in CHECKLIST:
        heading(group, size=13, colour="0A1D2B", space_before=10)
        for item in items:
            doc.add_paragraph(f"☐  {item}")

    heading("Twelve mistakes to check yourself against", size=15, colour="0A1D2B")
    for m in MISTAKES:
        doc.add_paragraph(m, style="List Bullet")

    heading("What still has to be decided", size=15, colour="0A1D2B")
    nt = doc.add_table(rows=1, cols=1)
    nt.style = "Table Grid"
    nt.rows[0].cells[0].text = "\n\n\n\n\n"

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = closing.add_run("You design the system. The agent runs it.")
    cr.bold = True
    cr.font.color.rgb = RGBColor.from_string("DD3C13")

    out = OUT.replace(".pdf", ".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    build(); print("wrote", OUT)
    print("wrote", build_docx())
